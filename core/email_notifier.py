import os
import smtplib
from email.mime.multipart import MIMEMultipart
from email.mime.text import MIMEText
from email.mime.application import MIMEApplication
from datetime import datetime
from config import (
    GMAIL_ADDRESS, GMAIL_APP_PASSWORD, NOTIFY_EMAIL,
    CV_PATH_RESEAUX, CV_PATH_AUTOMATISME, CV_PATH_JAVA, BRENDA_EMAIL,
)
from core.domain_classifier import DOMAIN_RESEAUX, DOMAIN_AUTOMATISME, DOMAIN_JAVA

DOMAIN_LABELS = {
    DOMAIN_RESEAUX: "Reseaux & Securite",
    DOMAIN_AUTOMATISME: "Automatisme & Systemes",
    DOMAIN_JAVA: "Java & Backend",
}

DOMAIN_COLORS = {
    DOMAIN_RESEAUX: "#1565C0",
    DOMAIN_AUTOMATISME: "#2E7D32",
    DOMAIN_JAVA: "#E65100",
}

CV_PATHS = {
    DOMAIN_RESEAUX: CV_PATH_RESEAUX,
    DOMAIN_AUTOMATISME: CV_PATH_AUTOMATISME,
    DOMAIN_JAVA: CV_PATH_JAVA,
}


def _score_badge(score: int) -> str:
    if score >= 75:
        return (f'<span style="background:#00C853;color:#000;padding:2px 8px;'
                f'border-radius:12px;font-weight:bold">{score}/100</span>')
    if score >= 50:
        return (f'<span style="background:#FFD600;color:#000;padding:2px 8px;'
                f'border-radius:12px;font-weight:bold">{score}/100</span>')
    return (f'<span style="background:#FF6D00;color:#fff;padding:2px 8px;'
            f'border-radius:12px;font-weight:bold">{score}/100</span>')


def _exp_badge(level: str) -> str:
    colors = {
        "Debutant": "#00C853", "Junior": "#40C4FF",
        "2-3 ans": "#FFD600", "3-5 ans": "#FF6D00",
        "Senior": "#F44336", "NC": "#9E9E9E",
    }
    color = colors.get(level, "#9E9E9E")
    return (f'<span style="background:{color};color:#000;padding:1px 6px;'
            f'border-radius:8px;font-size:11px">{level}</span>')


def _get_applied_urls(excel_paths: list = None) -> set:
    from core.excel_output import EXCEL_PATH_RESEAUX, EXCEL_PATH_AUTOMATISME, EXCEL_PATH_JAVA
    from openpyxl import load_workbook
    paths = excel_paths or [EXCEL_PATH_RESEAUX, EXCEL_PATH_AUTOMATISME, EXCEL_PATH_JAVA]
    applied = set()
    for path in paths:
        if not os.path.isfile(path):
            continue
        try:
            wb = load_workbook(path, read_only=True)
            ws = wb.active
            statut_col = 4
            url_col = 3
            for col in range(1, (ws.max_column or 20) + 1):
                header = ws.cell(1, col).value
                if header == "Statut":
                    statut_col = col
                elif header == "Lien de l'offre":
                    url_col = col
            for row in range(2, ws.max_row + 1):
                status = ws.cell(row, statut_col).value or ""
                if "postul" in str(status).lower():
                    url = ws.cell(row, url_col).value
                    if url:
                        applied.add(url)
        except Exception:
            pass
    return applied


def _skills_chips(skills: list, color: str) -> str:
    if not skills:
        return '<span style="color:#555;font-size:11px">—</span>'
    chips = "".join(
        f'<span style="background:{color}22;color:{color};border:1px solid {color}55;'
        f'padding:1px 7px;border-radius:10px;font-size:10px;margin:1px 2px 1px 0;'
        f'display:inline-block">{s}</span>'
        for s in skills[:6]
    )
    return chips


def _apply_btn(url: str, color: str, applied_urls: set, big: bool = False) -> str:
    if url in applied_urls:
        pad = "8px 18px" if big else "5px 12px"
        size = "13px" if big else "11px"
        return (
            f'<span style="background:#3949AB;color:#fff;padding:{pad};'
            f'border-radius:{"6" if big else "5"}px;font-size:{size};font-weight:bold;'
            f'display:inline-block;white-space:nowrap">&#10003; Deja postule</span>'
        )
    pad = "8px 18px" if big else "5px 12px"
    size = "13px" if big else "11px"
    arrow = " &rarr;" if big else ""
    return (
        f'<a href="{url}" style="background:{color};color:#fff;padding:{pad};'
        f'border-radius:{"6" if big else "5"}px;text-decoration:none;font-size:{size};'
        f'font-weight:bold;display:inline-block;white-space:nowrap">Postuler{arrow}</a>'
    )


def _top_priority_section(jobs: list, color: str, applied_urls: set) -> str:
    top = [j for j in jobs if j.relevance_score >= 80]
    if not top:
        return ""

    cards = ""
    for job in top[:5]:
        skills_html = _skills_chips(job.skills_detected, color)
        salary_txt = job.salary if job.salary and job.salary != "NC" else "Salaire NC"
        remote_txt = job.remote if job.remote != "NC" else ""
        remote_html = f' &nbsp;·&nbsp; {remote_txt}' if remote_txt and remote_txt != "Aucun" else ""
        btn = _apply_btn(job.url, color, applied_urls, big=True)
        cards += f"""
        <div style="background:#0A2540;border-radius:10px;padding:16px 20px;margin-bottom:12px;
                    border-left:4px solid {color}">
          <div style="display:flex;justify-content:space-between;align-items:flex-start;
                      flex-wrap:wrap;gap:8px">
            <div style="flex:1;min-width:200px">
              <div style="font-size:15px;font-weight:bold;color:#E0E0E0;margin-bottom:4px">
                {job.title}
              </div>
              <div style="font-size:13px;color:{color};margin-bottom:6px">
                {job.company} &nbsp;·&nbsp; {job.location}{remote_html}
              </div>
              <div style="font-size:12px;color:#9E9E9E;margin-bottom:8px">
                {salary_txt} &nbsp;·&nbsp; {_exp_badge(job.experience_level)}
              </div>
              <div style="margin-top:6px">{skills_html}</div>
            </div>
            <div style="display:flex;flex-direction:column;align-items:flex-end;gap:8px">
              {_score_badge(job.relevance_score)}
              {btn}
            </div>
          </div>
        </div>
        """

    return f"""
    <div style="background:#0D1B2A;border-radius:12px;padding:20px;margin-bottom:20px;
                border:1px solid {color}44">
      <div style="color:{color};font-size:14px;font-weight:bold;margin-bottom:14px;
                  letter-spacing:0.5px">
        &#9733; TOP PRIORITE &mdash; Score &ge; 80
      </div>
      {cards}
    </div>
    """


def _build_html(jobs: list, run_at: datetime, domain: str,
                applied_urls: set = None, recipient_name: str = "Stephane NANDJOU TONLEU") -> str:
    applied_urls = applied_urls or set()
    label = DOMAIN_LABELS.get(domain, domain)
    color = DOMAIN_COLORS.get(domain, "#1565C0")

    rows = ""
    for job in jobs:
        skills_html = _skills_chips(job.skills_detected, color)
        btn = _apply_btn(job.url, color, applied_urls, big=False)
        rows += f"""
        <tr style="border-bottom:1px solid #1E3A5F;">
          <td style="padding:10px 8px;color:#E0E0E0;font-size:13px">
            <a href="{job.url}"
               style="color:{color};text-decoration:none;font-weight:bold">{job.title}</a>
            <div style="font-size:11px;color:#9E9E9E;margin-top:2px">{job.company}</div>
            <div style="margin-top:5px">{skills_html}</div>
          </td>
          <td style="padding:10px 8px;color:#E0E0E0;font-size:12px">{job.source}</td>
          <td style="padding:10px 8px;color:#E0E0E0;font-size:12px">{job.location}</td>
          <td style="padding:10px 8px;color:#E0E0E0;font-size:12px">{job.salary}</td>
          <td style="padding:10px 8px;color:#E0E0E0;font-size:12px">{job.remote}</td>
          <td style="padding:10px 8px;font-size:12px">{_exp_badge(job.experience_level)}</td>
          <td style="padding:10px 8px;font-size:12px;text-align:center">
            {_score_badge(job.relevance_score)}
          </td>
          <td style="padding:10px 8px;color:#B0B0B0;font-size:11px">{job.age_hours}h</td>
          <td style="padding:10px 8px;text-align:center">{btn}</td>
        </tr>
        """

    sources_summary = {}
    for job in jobs:
        sources_summary[job.source] = sources_summary.get(job.source, 0) + 1
    sources_text = " · ".join(f"{src}: {count}" for src, count in sources_summary.items())

    top_section = _top_priority_section(jobs, color, applied_urls)

    return f"""
    <!DOCTYPE html>
    <html>
    <head><meta charset="UTF-8"></head>
    <body style="background:#081226;font-family:Calibri,Arial,sans-serif;margin:0;padding:20px">
      <div style="max-width:1050px;margin:0 auto">
        <div style="background:#0D1B2A;border-radius:12px;padding:24px;margin-bottom:20px;
                    border-left:4px solid {color}">
          <h1 style="color:{color};margin:0;font-size:22px">Job Agent - {label}</h1>
          <p style="color:#9E9E9E;margin:8px 0 0;font-size:13px">
            {run_at.strftime("%d/%m/%Y a %H:%M")} ·
            <strong style="color:#E0E0E0">{len(jobs)} nouvelle(s) offre(s)</strong> · {sources_text}
          </p>
        </div>
        {top_section}
        <div style="background:#0D1B2A;border-radius:12px;overflow:hidden">
          <table style="width:100%;border-collapse:collapse">
            <thead>
              <tr style="background:#0A2540">
                <th style="padding:12px 8px;color:{color};text-align:left;font-size:12px">Poste / Entreprise / Competences</th>
                <th style="padding:12px 8px;color:{color};text-align:left;font-size:12px">Source</th>
                <th style="padding:12px 8px;color:{color};text-align:left;font-size:12px">Lieu</th>
                <th style="padding:12px 8px;color:{color};text-align:left;font-size:12px">Salaire</th>
                <th style="padding:12px 8px;color:{color};text-align:left;font-size:12px">Teletravail</th>
                <th style="padding:12px 8px;color:{color};text-align:left;font-size:12px">Niveau</th>
                <th style="padding:12px 8px;color:{color};text-align:center;font-size:12px">Score</th>
                <th style="padding:12px 8px;color:{color};text-align:left;font-size:12px">Age</th>
                <th style="padding:12px 8px;color:{color};text-align:center;font-size:12px">Action</th>
              </tr>
            </thead>
            <tbody>{rows}</tbody>
          </table>
        </div>
        <p style="color:#4A5568;font-size:11px;text-align:center;margin-top:16px">
          Job Agent · {recipient_name}
        </p>
      </div>
    </body>
    </html>
    """


def _attach_file(msg, path: str, filename: str = None):
    if path and os.path.isfile(path):
        with open(path, "rb") as f:
            data = f.read()
        ext = os.path.splitext(path)[1].lower()
        subtype = "vnd.openxmlformats-officedocument.spreadsheetml.sheet" if ext == ".xlsx" else "pdf"
        part = MIMEApplication(data, _subtype=subtype)
        part.add_header("Content-Disposition", "attachment",
                        filename=filename or os.path.basename(path))
        msg.attach(part)
        print(f"[Email] PJ : {filename or os.path.basename(path)}")
        return True
    print(f"[Email] Fichier introuvable — PJ ignoree ({path})")
    return False


def send_alert(jobs: list, run_at: datetime = None):
    if not GMAIL_ADDRESS or not GMAIL_APP_PASSWORD:
        print("[Email] Credentials Gmail manquants — email ignore")
        return
    if not jobs:
        print("[Email] Aucune nouvelle offre — email non envoye")
        return

    run_at = run_at or datetime.now()
    applied_urls = _get_applied_urls([])
    if applied_urls:
        print(f"[Email] {len(applied_urls)} offre(s) deja postulees — marquees dans l'email")

    reseaux_jobs = [j for j in jobs if j.domain == DOMAIN_RESEAUX]
    auto_jobs = [j for j in jobs if j.domain == DOMAIN_AUTOMATISME]

    for domain, domain_jobs in [(DOMAIN_RESEAUX, reseaux_jobs), (DOMAIN_AUTOMATISME, auto_jobs)]:
        if not domain_jobs:
            continue
        label = DOMAIN_LABELS[domain]
        msg = MIMEMultipart("mixed")
        msg["Subject"] = (
            f"[Job Agent - {label}] {len(domain_jobs)} offre(s)"
            f" — {run_at.strftime('%d/%m/%Y %H:%M')}"
        )
        msg["From"] = GMAIL_ADDRESS
        msg["To"] = NOTIFY_EMAIL

        alt = MIMEMultipart("alternative")
        plain = f"{len(domain_jobs)} offre(s) {label} le {run_at.strftime('%d/%m/%Y a %H:%M')}.\n\n"
        for job in domain_jobs:
            applied_tag = " [DEJA POSTULE]" if job.url in applied_urls else ""
            plain += (
                f"- {job.title} — {job.company} ({job.location}){applied_tag}\n"
                f"  {job.url}\n  Score: {job.relevance_score}/100\n\n"
            )
        alt.attach(MIMEText(plain, "plain"))
        alt.attach(MIMEText(
            _build_html(domain_jobs, run_at, domain, applied_urls,
                        recipient_name="Stephane NANDJOU TONLEU · Ingenieur Reseaux &amp; Securite / Automaticien"),
            "html",
        ))
        msg.attach(alt)
        _attach_file(msg, CV_PATHS.get(domain, ""))

        try:
            with smtplib.SMTP_SSL("smtp.gmail.com", 465) as server:
                server.login(GMAIL_ADDRESS, GMAIL_APP_PASSWORD)
                server.sendmail(GMAIL_ADDRESS, NOTIFY_EMAIL, msg.as_string())
            print(f"[Email] OK {label} : {len(domain_jobs)} offres envoyees a {NOTIFY_EMAIL}")
        except Exception as e:
            print(f"[Email] ERREUR {label} : {e}")


def send_alert_brenda(jobs: list, run_at: datetime = None):
    from core.excel_output import EXCEL_PATH_JAVA

    if not GMAIL_ADDRESS or not GMAIL_APP_PASSWORD:
        print("[Email Brenda] Credentials Gmail manquants — email ignore")
        return
    if not jobs:
        print("[Email Brenda] Aucune nouvelle offre Java — email non envoye")
        return

    run_at = run_at or datetime.now()
    applied_urls = _get_applied_urls([EXCEL_PATH_JAVA])
    if applied_urls:
        print(f"[Email Brenda] {len(applied_urls)} offre(s) deja postulees")

    label = DOMAIN_LABELS[DOMAIN_JAVA]
    msg = MIMEMultipart("mixed")
    msg["Subject"] = (
        f"[Job Agent - {label}] {len(jobs)} offre(s)"
        f" — {run_at.strftime('%d/%m/%Y %H:%M')}"
    )
    msg["From"] = GMAIL_ADDRESS
    msg["To"] = BRENDA_EMAIL

    alt = MIMEMultipart("alternative")
    plain = f"Bonjour Brenda,\n\n{len(jobs)} nouvelle(s) offre(s) Java detectee(s) le {run_at.strftime('%d/%m/%Y a %H:%M')}.\n\n"
    for job in jobs:
        applied_tag = " [DEJA POSTULE]" if job.url in applied_urls else ""
        plain += (
            f"- {job.title} — {job.company} ({job.location}){applied_tag}\n"
            f"  {job.url}\n  Score: {job.relevance_score}/100\n\n"
        )
    alt.attach(MIMEText(plain, "plain"))
    alt.attach(MIMEText(
        _build_html(jobs, run_at, DOMAIN_JAVA, applied_urls,
                    recipient_name="Brenda KOUDJA · Ingenieure Logiciel Java"),
        "html",
    ))
    msg.attach(alt)

    # Excel Java en pièce jointe
    _attach_file(msg, EXCEL_PATH_JAVA, filename="jobs_java.xlsx")
    # CV de Brenda en pièce jointe
    _attach_file(msg, CV_PATH_JAVA)

    try:
        with smtplib.SMTP_SSL("smtp.gmail.com", 465) as server:
            server.login(GMAIL_ADDRESS, GMAIL_APP_PASSWORD)
            server.sendmail(GMAIL_ADDRESS, BRENDA_EMAIL, msg.as_string())
        print(f"[Email Brenda] OK : {len(jobs)} offres envoyees a {BRENDA_EMAIL}")
    except Exception as e:
        print(f"[Email Brenda] ERREUR : {e}")
