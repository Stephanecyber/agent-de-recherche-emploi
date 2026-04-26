"""
Générateur de lettres de motivation.
Format et style identiques aux vraies lettres de l'utilisateur.
Corps basé sur les lettres humaines existantes — pas de texte généré.
"""
import os
from datetime import datetime
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from core.domain_classifier import DOMAIN_RESEAUX

DATA_DIR = os.path.join(os.path.dirname(__file__), "..", "data", "lettres")

_MOIS = [
    "janvier", "février", "mars", "avril", "mai", "juin",
    "juillet", "août", "septembre", "octobre", "novembre", "décembre",
]

CANDIDAT = {
    "nom": "Stephane NANDJOU TONLEU",
    "adresse": "8 Allée Hélène Lazareff 78430 Louveciennes",
    "tel": "+33 6 18 80 78 09",
    "email": "tonleustephane1@gmail.com",
    "ville": "Louveciennes",
}


# ---------------------------------------------------------------------------
# Corps des lettres — basé sur les vraies lettres humaines de l'utilisateur
# Chaque paragraphe = liste de tuples (texte, gras)
# ---------------------------------------------------------------------------

def _body_reseaux(title: str, company: str, desc_lower: str) -> list:
    # Légère adaptation P3 selon ce que l'offre mentionne
    if any(x in desc_lower for x in ["siem", "soc", "threat"]):
        p3_adapt = (
            "Cette démarche me permet d'approfondir ma compréhension des vulnérabilités, "
            "de travailler sur des cas concrets d'analyse SIEM et de réponse aux incidents, "
            "et d'adopter une vision globale et cohérente de la sécurité des infrastructures."
        )
    elif any(x in desc_lower for x in ["azure", "cloud"]):
        p3_adapt = (
            "Cette démarche me permet d'approfondir ma compréhension des vulnérabilités, "
            "d'analyser des architectures réseau, système et cloud Azure, "
            "et de renforcer ma maîtrise des environnements hybrides et des bonnes pratiques de sécurité."
        )
    elif any(x in desc_lower for x in ["pentest", "audit", "nmap"]):
        p3_adapt = (
            "Cette démarche me permet d'approfondir ma compréhension des vulnérabilités "
            "à travers des exercices de tests d'intrusion, d'analyser des architectures réseau et système, "
            "et d'adopter une vision globale et cohérente de la sécurité des infrastructures."
        )
    else:
        p3_adapt = (
            "Cette démarche me permet d'approfondir ma compréhension des vulnérabilités, "
            "d'analyser des architectures réseau, système et cloud, "
            "et d'adopter une vision globale et cohérente de la sécurité des infrastructures."
        )

    return [
        [   # Paragraphe 1 — accroche
            ("Ingénieur réseaux et sécurité junior diplômé de l'ECAM-EPMI en ", False),
            ("Réseaux et Systèmes Intelligents", True),
            (", je souhaite vous proposer ma candidature pour le poste de ", False),
            (title, True),
            (f" au sein de {company}. "
             "Mon parcours m'a permis de développer une approche concrète et opérationnelle "
             "de la sécurité des systèmes d'information, orientée vers la maîtrise des risques, "
             "la fiabilité des environnements et la sécurité opérationnelle.", False),
        ],
        [   # Paragraphe 2 — stage BASSETTI avec métriques réelles
            ("Lors de mon stage de fin d'études chez ", False),
            ("BASSETTI Group", True),
            (", j'ai contribué à la sécurisation d'environnements critiques pour des clients "
             "grands comptes tels que ", False),
            ("INEOS", True), (", ", False), ("Hermès", True), (" et la ", False), ("RATP", True),
            (", notamment sur des problématiques de gestion des identités et des accès (IAM/RBAC), "
             "de sécurité applicative et de maintien en conditions opérationnelles. "
             "Ces interventions ont permis, entre autres, la sécurisation des accès pour ", False),
            ("plus de 500 utilisateurs", True),
            (" et la réduction de ", False),
            ("près de 95 % des risques", True),
            (" liés aux failles d'exportation de données sensibles, "
             "dans des contextes à fortes contraintes métier et techniques.", False),
        ],
        [   # Paragraphe 3 — démarche personnelle (adapté selon offre)
            ("Cette expérience m'a permis de travailler en étroite collaboration avec des équipes IT "
             "pluridisciplinaires, de comprendre les exigences de la sécurité en production et de "
             "développer une rigueur essentielle dans des environnements où la fiabilité et la "
             "continuité de service sont primordiales.\n\n"
             "En parallèle de mon parcours académique et professionnel, je m'investis activement "
             "dans le développement continu de mes compétences, notamment à travers un laboratoire "
             "de cybersécurité personnel et des plateformes de formation spécialisées. " + p3_adapt, False),
        ],
        [   # Paragraphe 4 — motivation / objectif
            ("Aujourd'hui, je recherche un poste en CDI au sein d'un environnement technique "
             "stimulant, dans lequel je pourrai apporter une contribution concrète, fiable et "
             "mesurable, tout en poursuivant le développement de mon expertise aux côtés "
             "d'équipes expérimentées.", False),
        ],
        [   # Paragraphe 5 — disponibilité
            ("Je serais ravi de pouvoir échanger avec vous afin de vous présenter plus en "
             "détail mon parcours et ma motivation.", False),
        ],
    ]


def _body_automatisme(title: str, company: str, desc_lower: str) -> list:
    # Adaptation P3 selon technologies détectées
    if any(x in desc_lower for x in ["scada", "supervision"]):
        p3_systems = (
            "J'ai notamment conçu des "
        )
        p3_bold = "systèmes de surveillance et de diagnostic d'équipements critiques"
        p3_end = (
            ", intégrant des notions de criticité, de modes dégradés et de validation "
            "fonctionnelle, ainsi que des systèmes de supervision SCADA. "
            "Ces projets m'ont permis de développer une vision globale des systèmes, "
            "depuis la définition des exigences jusqu'aux "
        )
        p3_bold2 = "tests et à la validation."
    elif any(x in desc_lower for x in ["rams", "amdec", "fmea", "sûreté", "surete"]):
        p3_systems = "J'ai notamment réalisé des analyses de "
        p3_bold = "sûreté de fonctionnement (RAMS, AMDEC)"
        p3_end = (
            " sur des systèmes à fortes contraintes de fiabilité, en appliquant des méthodes "
            "structurées depuis la définition des exigences jusqu'à la "
        )
        p3_bold2 = "validation et la vérification."
    else:
        p3_systems = "J'ai notamment conçu des "
        p3_bold = "systèmes de surveillance et de diagnostic d'équipements critiques"
        p3_end = (
            ", intégrant des notions de criticité, de modes dégradés et de validation "
            "fonctionnelle, ainsi que des systèmes automatisés sécurisés. "
            "Ces projets m'ont permis de développer une vision globale des systèmes, "
            "depuis la définition des exigences jusqu'aux "
        )
        p3_bold2 = "tests et à la validation."

    return [
        [   # Paragraphe 1 — accroche
            ("Ingénieur systèmes junior diplômé de l'ECAM-EPMI en ", False),
            ("Réseaux et Systèmes Intelligents", True),
            (", je souhaite vous proposer ma candidature pour le poste de ", False),
            (title, True),
            (f" au sein de {company}. "
             "Mon parcours m'a permis de construire une approche rigoureuse de l'ingénierie "
             "des systèmes, orientée analyse fonctionnelle, gestion des exigences, "
             "sûreté de fonctionnement et validation.", False),
        ],
        [   # Paragraphe 2 — stage BASSETTI
            ("Lors de mon stage de fin d'études chez ", False),
            ("BASSETTI Group", True),
            (", j'ai participé à des projets de ", False),
            ("digitalisation industrielle", True),
            (" en contribuant à la structuration de systèmes complexes, à la formalisation "
             "des règles métier et à la convergence entre environnements industriels (OT) et "
             "outils numériques (IT). Cette expérience m'a permis de comprendre l'importance "
             "de la cohérence des données, de la traçabilité et du travail collaboratif "
             "dans des contextes industriels exigeants.", False),
        ],
        [   # Paragraphe 3 — projets ECAM (adapté)
            ("En parallèle, mes projets d'ingénierie réalisés à l'ECAM-EPMI m'ont amené à "
             "travailler sur des systèmes à fortes contraintes de fiabilité et de sécurité. "
             + p3_systems, False),
            (p3_bold, True),
            (p3_end, False),
            (p3_bold2, True),
        ],
        [   # Paragraphe 4 — soft skills
            ("Au-delà des compétences techniques, je me définis comme un ingénieur curieux, "
             "rigoureux et engagé, appréciant le travail en équipe et le dialogue entre les "
             "métiers. Je suis particulièrement motivé par les environnements où la fiabilité, "
             "la sécurité et la qualité des solutions techniques sont des enjeux majeurs.", False),
        ],
        [   # Paragraphe 5 — disponibilité
            ("Disponible immédiatement, je serais ravi de pouvoir échanger avec vous afin de "
             "vous présenter plus en détail mon parcours, mes projets et ma motivation à "
             "m'investir durablement au sein de vos équipes.", False),
        ],
    ]


# ---------------------------------------------------------------------------
# Helpers docx
# ---------------------------------------------------------------------------

def _font(run, size: int, bold: bool = False, italic: bool = False,
          color: RGBColor = None, name: str = "Calibri"):
    run.font.name = name
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    if color:
        run.font.color.rgb = color


def _para(doc, align=WD_ALIGN_PARAGRAPH.LEFT, space_before=0, space_after=0) -> object:
    p = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)
    return p


def _remove_table_borders(table):
    tbl = table._tbl
    tblPr = tbl.tblPr
    if tblPr is None:
        tblPr = OxmlElement("w:tblPr")
        tbl.insert(0, tblPr)
    tblBorders = OxmlElement("w:tblBorders")
    for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{side}")
        el.set(qn("w:val"), "none")
        el.set(qn("w:sz"), "0")
        tblBorders.append(el)
    tblPr.append(tblBorders)


def _add_mixed_para(doc, segments: list, size: int = 11,
                    align=WD_ALIGN_PARAGRAPH.JUSTIFY,
                    space_before: int = 0, space_after: int = 8):
    """Ajoute un paragraphe avec des runs de formatage mixte (gras/normal)."""
    p = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)
    for text, bold in segments:
        run = p.add_run(text)
        _font(run, size, bold=bold)
    return p


# ---------------------------------------------------------------------------
# Construction du document Word
# ---------------------------------------------------------------------------

def _build_docx(job, save_dir: str) -> str:
    domain = getattr(job, "domain", DOMAIN_RESEAUX)
    is_reseaux = domain == DOMAIN_RESEAUX
    company = job.company if job.company not in ("Entreprise NC", "") else "votre entreprise"
    desc_lower = job.description.lower()
    now_dt = datetime.now()
    date_str = f"{now_dt.day} {_MOIS[now_dt.month - 1]} {now_dt.year}"

    title_label = (
        "Ingénieur Réseaux et Sécurité / Cybersécurité Junior"
        if is_reseaux else
        "Ingénieur Systèmes Junior - Systèmes Critiques"
    )
    objet = (
        f"Candidature pour le poste de {job.title}"
        if job.title else
        "Candidature Ingénieur Junior"
    )

    body = (
        _body_reseaux(job.title, company, desc_lower)
        if is_reseaux else
        _body_automatisme(job.title, company, desc_lower)
    )

    doc = Document()

    # Marges identiques à la vraie lettre
    for section in doc.sections:
        section.top_margin = Cm(2.2)
        section.bottom_margin = Cm(2.2)
        section.left_margin = Cm(2.5)
        section.right_margin = Cm(2.5)

    # -----------------------------------------------------------------------
    # EN-TÊTE : tableau 2 colonnes (contact gauche | destination droite)
    # -----------------------------------------------------------------------
    table = doc.add_table(rows=1, cols=2)
    _remove_table_borders(table)
    table.columns[0].width = Cm(9)
    table.columns[1].width = Cm(7)

    left = table.cell(0, 0)
    right = table.cell(0, 1)

    # Colonne gauche — coordonnées
    def _lp(cell, text, bold=False, italic=False, size=10.5, after=1):
        p = cell.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_after = Pt(after)
        r = p.add_run(text)
        _font(r, size, bold=bold, italic=italic)
        return p

    left.paragraphs[0].clear()
    _lp(left, f"Nom : {CANDIDAT['nom']}", bold=True, size=11, after=2)
    _lp(left, f"Adresse : {CANDIDAT['adresse']}", after=2)
    _lp(left, f"Tel : {CANDIDAT['tel']}", after=2)
    _lp(left, f"Email : {CANDIDAT['email']}", after=4)
    _lp(left, title_label, italic=True, after=0)

    # Colonne droite — destinataire + date
    def _rp(cell, text, size=10.5, after=2):
        p = cell.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        p.paragraph_format.space_after = Pt(after)
        r = p.add_run(text)
        _font(r, size)
        return p

    right.paragraphs[0].clear()
    _rp(right, "A l'attention du Responsable du Recrutement")
    _rp(right, "Service des Ressources Humaines")
    _rp(right, f"À {CANDIDAT['ville']}, le {date_str}", after=0)

    # Espace après tableau
    doc.add_paragraph().paragraph_format.space_after = Pt(6)

    # -----------------------------------------------------------------------
    # OBJET
    # -----------------------------------------------------------------------
    p_objet = _para(doc, space_after=10)
    r = p_objet.add_run(f"Objet : {objet}")
    _font(r, 11, bold=True)

    # -----------------------------------------------------------------------
    # FORMULE D'APPEL
    # -----------------------------------------------------------------------
    p_appel = _para(doc, space_after=10)
    r = p_appel.add_run("Madame, Monsieur,")
    _font(r, 11, bold=True)

    # -----------------------------------------------------------------------
    # CORPS DE LA LETTRE
    # -----------------------------------------------------------------------
    for i, segments in enumerate(body):
        space_after = 8 if i < len(body) - 1 else 12
        _add_mixed_para(doc, segments, space_after=space_after)

    # -----------------------------------------------------------------------
    # FORMULE DE POLITESSE
    # -----------------------------------------------------------------------
    p_pol = _para(doc, space_after=20)
    r = p_pol.add_run(
        "Je vous prie d'agréer, Madame, Monsieur, l'expression de mes salutations distinguées."
    )
    _font(r, 11)

    # Espace signature
    _para(doc, space_after=8)

    # Signature — alignée à droite comme dans les vraies lettres
    p_sig = _para(doc, align=WD_ALIGN_PARAGRAPH.RIGHT, space_after=2)
    r = p_sig.add_run(CANDIDAT["nom"])
    _font(r, 11, bold=True)

    p_titre = _para(doc, align=WD_ALIGN_PARAGRAPH.RIGHT, space_after=0)
    r = p_titre.add_run(title_label)
    _font(r, 10.5, italic=True)

    # -----------------------------------------------------------------------
    # Sauvegarde
    # -----------------------------------------------------------------------
    safe_co = "".join(c if c.isalnum() or c in "-_" else "_" for c in company)[:28]
    safe_ti = "".join(c if c.isalnum() or c in "-_" else "_" for c in job.title)[:28]
    docx_path = os.path.join(save_dir, f"LM_{safe_co}_{safe_ti}.docx")
    doc.save(docx_path)
    return docx_path


def _to_pdf(docx_path: str) -> str:
    try:
        from docx2pdf import convert
        pdf_path = docx_path.replace(".docx", ".pdf")
        convert(docx_path, pdf_path)
        return pdf_path
    except Exception as e:
        print(f"[LM] PDF non généré ({e}) — le .docx est disponible")
        return ""


# ---------------------------------------------------------------------------
# API publique
# ---------------------------------------------------------------------------

def generate(job, output_dir: str = None) -> tuple:
    """Génère Word + PDF. Retourne (docx_path, pdf_path)."""
    save_dir = output_dir or DATA_DIR
    os.makedirs(save_dir, exist_ok=True)
    docx_path = _build_docx(job, save_dir)
    pdf_path = _to_pdf(docx_path)
    return docx_path, pdf_path


def generate_batch(jobs: list, top_n: int = 10) -> list:
    """Génère lettres pour les N meilleures offres (score >= 60)."""
    top_jobs = sorted(
        [j for j in jobs if j.relevance_score >= 60],
        key=lambda j: (-j.relevance_score, j.age_hours),
    )[:top_n]

    results = []
    for job in top_jobs:
        try:
            docx_path, pdf_path = generate(job)
            results.append((job, docx_path, pdf_path))
        except Exception as e:
            print(f"[LM] ERREUR {job.title} : {e}")
    return results
