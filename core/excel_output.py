import os
from datetime import datetime
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.formatting.rule import FormulaRule
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from core.domain_classifier import DOMAIN_RESEAUX, DOMAIN_AUTOMATISME

DATA_DIR = os.path.join(os.path.dirname(__file__), "..", "data")
EXCEL_PATH_RESEAUX = os.path.join(DATA_DIR, "jobs_reseaux_secu.xlsx")
EXCEL_PATH_AUTOMATISME = os.path.join(DATA_DIR, "jobs_automatisme.xlsx")

# Colonnes ordonnées : les plus utiles d'abord (visibles sans scroll)
COLUMNS = [
    ("Titre du poste", 38),        # 1
    ("Entreprise", 26),             # 2
    ("Lien de l'offre", 42),        # 3
    ("Statut", 14),                 # 4
    ("Date publication", 20),       # 5
    ("Niveau experience", 18),      # 6  ← rapproché de Date publication
    ("Score /100", 12),             # 7
    ("Source", 18),                 # 8
    ("Lieu", 24),                   # 9
    ("Region", 18),                 # 10
    ("Salaire", 24),                # 11
    ("Teletravail", 15),            # 12
    ("Contrat", 10),                # 13
    ("Age (h)", 10),                # 14
    ("Resume description", 60),     # 15
    ("Competences detectees", 36),  # 16
    ("Notes personnelles", 30),     # 17
    ("Detecte le", 20),             # 18
]

# Indice des colonnes clés (1-based)
COL_STATUT = 4
COL_DATE_PUB = 5
COL_NIVEAU = 6
COL_SCORE = 7
COL_LIEN = 3
COL_NOTES = 17
COL_DETECTE = 18

HEADER_FILL_RESEAUX = PatternFill("solid", fgColor="1565C0")
HEADER_FILL_AUTOMATISME = PatternFill("solid", fgColor="2E7D32")
HEADER_FONT = Font(bold=True, color="FFFFFF", size=10, name="Calibri")

ROW_FILL_ODD = PatternFill("solid", fgColor="FFFFFF")
ROW_FILL_EVEN_RESEAUX = PatternFill("solid", fgColor="EEF4FF")
ROW_FILL_EVEN_AUTOMATISME = PatternFill("solid", fgColor="E8F5E9")

SCORE_HIGH_FILL = PatternFill("solid", fgColor="C8F5D8")
SCORE_MED_FILL = PatternFill("solid", fgColor="FFF9C4")
SCORE_LOW_FILL = PatternFill("solid", fgColor="FFE0B2")

BORDER_COLOR = "BBCDE5"
THIN = Border(
    left=Side(style="thin", color=BORDER_COLOR),
    right=Side(style="thin", color=BORDER_COLOR),
    top=Side(style="thin", color=BORDER_COLOR),
    bottom=Side(style="thin", color=BORDER_COLOR),
)

TEXT_COLOR = "1A1A2E"
LINK_COLOR_RESEAUX = "1565C0"
LINK_COLOR_AUTOMATISME = "2E7D32"
MUTED_COLOR = "546E7A"
NEW_STATUS_COLOR = "1B5E20"

# Valeurs autorisées dans le dropdown Statut
STATUS_VALUES = "Nouveau,Postulé,Entretien,Refusé,Abandonné"

# Nombre max de lignes couvertes par la validation et le formatage conditionnel
MAX_ROWS = 2000


def _score_fill(score: int) -> PatternFill:
    if score >= 75:
        return SCORE_HIGH_FILL
    if score >= 50:
        return SCORE_MED_FILL
    return SCORE_LOW_FILL


def _score_font_color(score: int) -> str:
    if score >= 75:
        return "1B5E20"
    if score >= 50:
        return "F57F17"
    return "BF360C"


def _add_status_validation_and_cf(ws):
    """Ajoute le dropdown Statut et le formatage conditionnel sur les lignes."""
    col_letter = get_column_letter(COL_STATUT)
    data_range = f"{col_letter}2:{col_letter}{MAX_ROWS}"
    row_range = f"A2:{get_column_letter(len(COLUMNS))}{MAX_ROWS}"

    # Dropdown
    dv = DataValidation(
        type="list",
        formula1=f'"{STATUS_VALUES}"',
        allow_blank=True,
        showDropDown=False,
        showErrorMessage=False,
    )
    ws.add_data_validation(dv)
    dv.sqref = data_range

    # Ligne "Postulé" → violet/indigo
    ws.conditional_formatting.add(
        row_range,
        FormulaRule(
            formula=[f'${col_letter}2="Postulé"'],
            fill=PatternFill("solid", fgColor="E8EAF6"),
            font=Font(color="3949AB", name="Calibri", size=10),
            stopIfTrue=True,
        ),
    )

    # Ligne "Entretien" → cyan clair
    ws.conditional_formatting.add(
        row_range,
        FormulaRule(
            formula=[f'${col_letter}2="Entretien"'],
            fill=PatternFill("solid", fgColor="E0F7FA"),
            font=Font(color="00695C", name="Calibri", size=10),
            stopIfTrue=True,
        ),
    )

    # Ligne "Refusé" → rouge pâle
    ws.conditional_formatting.add(
        row_range,
        FormulaRule(
            formula=[f'${col_letter}2="Refusé"'],
            fill=PatternFill("solid", fgColor="FFEAEA"),
            font=Font(color="C62828", name="Calibri", size=10),
            stopIfTrue=True,
        ),
    )

    # Ligne "Abandonné" → gris
    ws.conditional_formatting.add(
        row_range,
        FormulaRule(
            formula=[f'${col_letter}2="Abandonné"'],
            fill=PatternFill("solid", fgColor="F5F5F5"),
            font=Font(color="9E9E9E", name="Calibri", size=10),
        ),
    )


def _create_workbook(domain: str) -> Workbook:
    wb = Workbook()
    ws = wb.active
    if domain == DOMAIN_RESEAUX:
        ws.title = "Reseaux & Securite"
        ws.sheet_properties.tabColor = "1565C0"
        header_fill = HEADER_FILL_RESEAUX
    else:
        ws.title = "Automatisme & Systemes"
        ws.sheet_properties.tabColor = "2E7D32"
        header_fill = HEADER_FILL_AUTOMATISME

    ws.freeze_panes = "A2"

    for col_idx, (col_name, col_width) in enumerate(COLUMNS, start=1):
        cell = ws.cell(row=1, column=col_idx, value=col_name)
        cell.fill = header_fill
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = THIN
        ws.column_dimensions[get_column_letter(col_idx)].width = col_width

    ws.row_dimensions[1].height = 32
    ws.auto_filter.ref = f"A1:{get_column_letter(len(COLUMNS))}1"

    _add_status_validation_and_cf(ws)

    return wb


def _add_job_row(ws, job, row_idx: int, domain: str):
    is_reseaux = domain == DOMAIN_RESEAUX
    base_fill = ROW_FILL_ODD if row_idx % 2 != 0 else (
        ROW_FILL_EVEN_RESEAUX if is_reseaux else ROW_FILL_EVEN_AUTOMATISME
    )
    link_color = LINK_COLOR_RESEAUX if is_reseaux else LINK_COLOR_AUTOMATISME
    score = job.relevance_score

    pub_local = job.published_at.astimezone().replace(tzinfo=None) \
        if job.published_at.tzinfo else job.published_at
    detected_local = job.detected_at.astimezone().replace(tzinfo=None) \
        if job.detected_at.tzinfo else job.detected_at

    # Ordre des valeurs aligné sur COLUMNS
    values = [
        job.title,                                                     # 1  Titre
        job.company,                                                   # 2  Entreprise
        job.url,                                                       # 3  Lien
        job.status,                                                    # 4  Statut
        pub_local,                                                     # 5  Date publication
        job.experience_level,                                          # 6  Niveau experience
        score,                                                         # 7  Score
        job.source,                                                    # 8  Source
        job.location,                                                  # 9  Lieu
        job.region,                                                    # 10 Region
        job.salary,                                                    # 11 Salaire
        job.remote,                                                    # 12 Teletravail
        job.contract_type,                                             # 13 Contrat
        job.age_hours,                                                 # 14 Age (h)
        job.description,                                               # 15 Resume description
        ", ".join(job.skills_detected) if job.skills_detected else "", # 16 Competences
        job.notes,                                                     # 17 Notes
        detected_local,                                                # 18 Detecte le
    ]

    for col_idx, value in enumerate(values, start=1):
        cell = ws.cell(row=row_idx, column=col_idx, value=value)
        cell.border = THIN
        cell.alignment = Alignment(
            vertical="center",
            wrap_text=(col_idx in (1, 15, 16, 17)),
            horizontal="left",
        )

        if col_idx == COL_SCORE:  # 7
            cell.fill = _score_fill(score)
            cell.font = Font(bold=True, color=_score_font_color(score), name="Calibri", size=10)
            cell.alignment = Alignment(horizontal="center", vertical="center")
        elif col_idx == COL_LIEN:  # 3
            cell.hyperlink = value or ""
            cell.font = Font(color=link_color, underline="single", name="Calibri", size=10)
            cell.fill = base_fill
        elif col_idx == COL_STATUT:  # 4
            val_str = str(value or "")
            is_applied = "postul" in val_str.lower()
            is_new = "Nouveau" in val_str
            if is_applied:
                cell.fill = PatternFill("solid", fgColor="E8EAF6")
                cell.font = Font(color="3949AB", bold=True, name="Calibri", size=10)
            elif is_new:
                cell.fill = PatternFill("solid", fgColor="D4EDDA")
                cell.font = Font(color=NEW_STATUS_COLOR, bold=True, name="Calibri", size=10)
            else:
                cell.fill = base_fill
                cell.font = Font(color=TEXT_COLOR, name="Calibri", size=10)
        elif col_idx in (8, 12, 13, 14):  # Source, Teletravail, Contrat, Age
            cell.fill = base_fill
            cell.font = Font(color=MUTED_COLOR, name="Calibri", size=10)
            cell.alignment = Alignment(horizontal="center", vertical="center")
        elif col_idx == COL_NIVEAU:  # 6
            cell.fill = base_fill
            cell.font = Font(color=MUTED_COLOR, name="Calibri", size=10)
            cell.alignment = Alignment(horizontal="center", vertical="center")
        elif col_idx in (1, 2):
            cell.fill = base_fill
            cell.font = Font(color=TEXT_COLOR, bold=(col_idx == 1), name="Calibri", size=10)
        else:
            cell.fill = base_fill
            cell.font = Font(color=TEXT_COLOR, name="Calibri", size=10)

        if col_idx in (COL_DATE_PUB, COL_DETECTE):  # 5, 18
            cell.number_format = "DD/MM/YYYY HH:MM"

    ws.row_dimensions[row_idx].height = 38


def _detect_col_map(ws) -> dict:
    """Lit l'en-tête ligne 1 et retourne {nom_colonne: index_1based}.
    Permet de lire les anciens fichiers même si les colonnes ont été réordonnées."""
    return {
        ws.cell(1, col).value: col
        for col in range(1, ws.max_column + 1)
        if ws.cell(1, col).value
    }


def _read_existing_jobs(ws, domain: str) -> list:
    from core.models import Job

    hdr = _detect_col_map(ws)

    def c(name, fallback=None):
        return hdr.get(name, fallback)

    jobs = []
    for row_idx in range(2, ws.max_row + 1):
        def v(col):
            return ws.cell(row_idx, col).value if col else None

        title = v(c("Titre du poste", 1))
        if not title:
            continue
        try:
            pub = v(c("Date publication", COL_DATE_PUB))
            if isinstance(pub, str):
                pub = datetime.fromisoformat(pub)
            det_raw = v(c("Detecte le", COL_DETECTE))
            detected = datetime.fromisoformat(det_raw) if isinstance(det_raw, str) else det_raw
            skills_raw = v(c("Competences detectees", 15)) or ""
            skills = [s.strip() for s in skills_raw.split(",") if s.strip()]
            job = Job(
                id=str(row_idx),
                title=title or "",
                company=v(c("Entreprise", 2)) or "",
                url=v(c("Lien de l'offre", 3)) or "",
                source=v(c("Source", 7)) or "",
                location=v(c("Lieu", 8)) or "",
                region=v(c("Region", 9)) or "",
                salary=v(c("Salaire", 10)) or "NC",
                remote=v(c("Teletravail", 11)) or "NC",
                contract_type=v(c("Contrat", 12)) or "NC",
                published_at=pub or datetime.now(),
                age_hours=float(v(c("Age (h)", 13)) or 0),
                relevance_score=int(v(c("Score /100", COL_SCORE)) or 0),
                description=v(c("Resume description", 14)) or "",
                skills_detected=skills,
                experience_level=v(c("Niveau experience", 16)) or "NC",
                status=v(c("Statut", COL_STATUT)) or "Nouveau",
                notes=v(c("Notes personnelles", COL_NOTES)) or "",
                detected_at=detected or datetime.now(),
                domain=domain,
            )
            jobs.append(job)
        except Exception:
            pass
    return jobs


def _save_domain_jobs(jobs: list, path: str, domain: str) -> str:
    os.makedirs(DATA_DIR, exist_ok=True)

    existing = []
    if os.path.exists(path):
        try:
            wb_old = load_workbook(path)
            existing = _read_existing_jobs(wb_old.active, domain)
        except Exception:
            pass

    jobs_all = jobs + [j for j in existing if j.url not in {x.url for x in jobs}]

    def _to_naive(dt):
        if dt is None:
            return datetime.min
        return dt.astimezone().replace(tzinfo=None) if dt.tzinfo else dt

    jobs_all.sort(key=lambda j: _to_naive(j.published_at), reverse=True)

    wb = _create_workbook(domain)
    ws = wb.active
    for idx, job in enumerate(jobs_all, start=2):
        _add_job_row(ws, job, idx, domain)

    wb.save(path)
    return path


def save_jobs(jobs: list) -> tuple:
    reseaux_jobs = [j for j in jobs if j.domain == DOMAIN_RESEAUX]
    auto_jobs = [j for j in jobs if j.domain == DOMAIN_AUTOMATISME]

    paths = []
    if reseaux_jobs:
        path = _save_domain_jobs(reseaux_jobs, EXCEL_PATH_RESEAUX, DOMAIN_RESEAUX)
        print(f"Excel Reseaux & Securite : {path} ({len(reseaux_jobs)} offres)")
        paths.append(path)
    if auto_jobs:
        path = _save_domain_jobs(auto_jobs, EXCEL_PATH_AUTOMATISME, DOMAIN_AUTOMATISME)
        print(f"Excel Automatisme & Systemes : {path} ({len(auto_jobs)} offres)")
        paths.append(path)

    return tuple(paths)
